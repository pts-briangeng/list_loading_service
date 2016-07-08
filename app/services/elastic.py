import abc
import csv
import httplib
import json
import logging
import os
import shutil
import traceback

import elasticsearch
import openpyxl
from elasticsearch import helpers, exceptions
from elasticsearch.client.utils import query_params
from liblcp import context
from requestswrapper import requests_wrapper

import configuration

logger = logging.getLogger(__name__)


class BulkAccountsFileReaders(object):

    class FileReader(object):

        def __init__(self, filename):
            self.filename = filename
            self.descriptor = None

        @abc.abstractmethod
        def get_rows(self):
            pass

        @abc.abstractmethod
        def is_empty(self):
            pass

        @abc.abstractmethod
        def close(self):
            pass

    class CsvReader(FileReader):

        def __init__(self, filename):
            super(BulkAccountsFileReaders.CsvReader, self).__init__(filename)
            self.descriptor = open(filename, 'rU')

        def close(self):
            self.descriptor.close()

        def get_rows(self):
            for row in csv.reader(self.descriptor):
                yield row[0]

        def is_empty(self):
            return os.stat(self.filename).st_size == 0

    class ExcelReader(FileReader):

        def __init__(self, filename):
            super(BulkAccountsFileReaders.ExcelReader, self).__init__(filename)
            self.workbook = openpyxl.load_workbook(self.filename, read_only=True)
            self.descriptor = self.workbook.active

        def get_rows(self):
            for row in self.descriptor.rows:
                yield row[0].value

        def is_empty(self):
            return sum(1 for _ in self.descriptor.rows) == 0

        def close(self):
            pass

    @classmethod
    def get(cls, file_path):
        file_type = file_path.split('.')[-1].lower()
        if file_type in ['csv', 'txt']:
            return BulkAccountsFileReaders.CsvReader(file_path)
        return BulkAccountsFileReaders.ExcelReader(file_path)


def rename_file(request_file, request_list):
    def rreplace(s, old, new, occurrence):
        li = s.rsplit(old, occurrence)
        return new.join(li)

    file_name, _ = os.path.splitext(os.path.basename(request_file))
    return rreplace(request_file, file_name, request_list, 1)


def elastic_search_callback(f):
    def wrapper(request):
        errors = False
        try:
            f(request)
        except Exception as e:
            errors = True
            logger.error('An error occurred when creating a new list: {}'.format(e.message))
            logger.error(traceback.format_exc(), exc_info=1)
        finally:
            if request.callbackUrl:
                data = {
                    'success': not errors,
                    'file': rename_file(request.filePath, request.list_id),
                    'links': {
                        'self': {
                            'href': request.url
                        }
                    }
                }
                if not errors:
                    data['links']['member'] = {
                        'href': '/{}/{}/{{member-id}}'.format(request.service, request.list_id)
                    }

                requests_wrapper.post(url=request.callbackUrl, data=json.dumps(data),
                                      headers=dict(context.get_headers(), **{'Content-Type': 'application/json'}))

    return wrapper


class ElasticSearchClient(elasticsearch.Elasticsearch):

    def __init__(self, **kwargs):
        super(ElasticSearchClient, self).__init__(hosts=[configuration.data.ELASTIC_SEARCH_SERVER], **kwargs)

    def _create_es_index_if_required(self, index):
        try:
            if not self.indices.exists(index=index):
                logger.info("Creating new index {}".format(index))
                self.indices.create(index=index)
        except exceptions.TransportError as e:
            logger.warning("Elastic search get index request exception!")
            logger.warning(traceback.format_exc(), exc_info=1)
            raise e

    def _create_es_mapping(self, index, doc_type):
        try:
            self.indices.put_mapping(
                doc_type=doc_type,
                index=index,
                body={
                    "properties": {
                        "accountNumber": {
                            "type": "string",
                            "index": "not_analyzed"
                        }
                    }
                })
        except exceptions.TransportError as e:
            logger.warning("Elastic search create mapping request exception")
            logger.warning(traceback.format_exc(), exc_info=1)
            raise e

    @query_params('consistency', 'refresh', 'routing', 'replication', 'timeout')
    def bulk(self, body, index=None, doc_type=None, params=None):
        self._create_es_index_if_required(index)
        self._create_es_mapping(index, doc_type)
        return super(ElasticSearchClient, self).bulk(body, index=index, doc_type=doc_type, params=params)


class ElasticSearchDocument(object):

    def __init__(self, index, type, account_number):
        self.index = index
        self.type = type
        self.account_number = account_number

    @property
    def doc(self):
        return {
            "_index": self.index,
            "_type": self.type,
            "_id": self.account_number,
            "_source": {
                "accountNumber": self.account_number
            }
        }


class ElasticSearchService(object):

    @staticmethod
    @elastic_search_callback
    def create_list(request):
        file_path = os.path.join(configuration.data.VOLUME_MAPPINGS_FILE_UPLOAD_TARGET, request.filePath)
        if not os.path.isfile(file_path):
            raise IOError("File {} does not exist!".format(file_path))

        updated_path = rename_file(file_path, request.list_id)
        shutil.move(file_path, updated_path)

        file_reader = BulkAccountsFileReaders.get(updated_path)
        if file_reader.is_empty():
            raise EOFError("File {} is empty!".format(file_path))

        actions = (ElasticSearchDocument(index=request.service, type=request.list_id, account_number=line).doc
                   for line in file_reader.get_rows())

        logger.info("Bulk indexing file using index: {}, type: {}".format(request.service, request.list_id))
        elastic_search_client = ElasticSearchClient()
        result = helpers.bulk(elastic_search_client, actions, index=request.service, doc_type=request.list_id)
        logger.info("Uploading ...Done! Refresh index")
        elastic_search_client.indices.refresh(index=request.service)
        logger.info("Finished indexing {} documents".format(result[0]))
        file_reader.close()
        return updated_path

    @staticmethod
    def delete_list(request):
        file_path = os.path.join(configuration.data.VOLUME_MAPPINGS_FILE_UPLOAD_TARGET, request.filePath)
        try:
            logger.info("Elasticsearch is deleting index: {}, doc_type: {}".format(request.service, request.list_id))
            elastic_search_client = ElasticSearchClient()
            result = elastic_search_client.indices.delete_mapping(index=request.service, doc_type=request.list_id)
            logger.info("Elastic search delete response {}".format(result))
        except exceptions.TransportError as e:
            if e.status_code == httplib.NOT_FOUND:
                logger.warning("Elastic search delete request not found")
                raise LookupError
            else:
                logger.warning("Elastic search delete request exception: {}".format(e.info))
                raise e

        if not result.get('acknowledged', False):
            logger.warning("Elastic search delete response not acknowledged successfully")
            raise Exception

        try:
            os.remove(file_path)
        except OSError as e:
            logger.warning("Error deleting file: {}".format(e))

        return result

    @staticmethod
    def get_list_status(request):
        elastic_search_client = ElasticSearchClient()
        result = elastic_search_client.search(index=request.service, doc_type=request.list_id, search_type="count")
        logger.info("elastic search response {}".format(result))
        if result['hits']['total'] == 0:
            logger.warning("Elastic search (index:{}, Type:{}) not found!".format(request.service, request.list_id))
            raise LookupError
        return result

    @staticmethod
    def get_list_member(request):
        elastic_search_client = ElasticSearchClient()
        if not elastic_search_client.exists(index=request.service, doc_type=request.list_id, id=request.member_id):
            raise LookupError
        return {}
